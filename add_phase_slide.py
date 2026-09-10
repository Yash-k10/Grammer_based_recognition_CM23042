"""
Script to add the 3-Phase Execution & Completion Slide to Grammar-Based_Pattern_Recognition_TAE1.pptx
Inserts the slide right before the final 'THANK YOU' slide.
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def add_phase_completion_slide(pptx_path: str):
    prs = Presentation(pptx_path)

    # Color Palette from presentation
    BG_COLOR = RGBColor(0xF2, 0xF7, 0xF7)       # #F2F7F7
    PRIMARY_DARK = RGBColor(0x0B, 0x30, 0x33)   # #0B3033
    TEAL_ACCENT = RGBColor(0x02, 0x80, 0x90)    # #028090
    EMERALD_GREEN = RGBColor(0x02, 0xC3, 0x9A)  # #02C39A
    TEXT_MUTED = RGBColor(0x5C, 0x7A, 0x7D)     # #5C7A7D
    CARD_BG = RGBColor(0xFF, 0xFF, 0xFF)        # #FFFFFF
    CARD_BORDER = RGBColor(0xC8, 0xDC, 0xDC)    # #C8DCDC
    BADGE_BG_DONE = RGBColor(0xD7, 0xEA, 0xEA)  # #D7EAEA
    BADGE_TEXT_DONE = RGBColor(0x00, 0x7A, 0x6E)# #007A6E
    BADGE_BG_NEXT = RGBColor(0xEB, 0xF0, 0xF0)  # #EBF0F0
    BADGE_TEXT_NEXT = RGBColor(0x5C, 0x7A, 0x7D)# #5C7A7D
    HEADER_ACCENT_1 = RGBColor(0x02, 0x80, 0x90)
    HEADER_ACCENT_2 = RGBColor(0x00, 0xA8, 0x96)
    HEADER_ACCENT_3 = RGBColor(0x3B, 0x6E, 0x72)

    # If presentation already has 7 slides, remove old roadmap slide at index 5 for clean replacement
    if len(prs.slides) >= 7:
        rId = prs.slides._sldIdLst[5].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[5]

    # Add new blank slide
    blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(blank_layout)

    # Reorder slide so it sits at index 5 (before the final slide)
    sldIdLst = prs.slides._sldIdLst
    new_id = sldIdLst[-1]
    sldIdLst.remove(new_id)
    sldIdLst.insert(5, new_id)

    # Set background color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

    # 1. Header Tag
    tag_box = slide.shapes.add_textbox(548640, 411480, 4572000, 274320)
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "PROJECT ROADMAP"
    p_tag.font.name = "Arial"
    p_tag.font.size = Pt(10.0)
    p_tag.font.bold = True
    p_tag.font.color.rgb = TEAL_ACCENT

    # 2. Main Title
    title_box = slide.shapes.add_textbox(548640, 685800, 7315200, 548640)
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = "Phase-Wise Project Completion"
    p_title.font.name = "Arial"
    p_title.font.size = Pt(28.0)
    p_title.font.bold = True
    p_title.font.color.rgb = PRIMARY_DARK

    # 3. Subtitle
    sub_box = slide.shapes.add_textbox(548640, 1234440, 11094720, 320000)
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Structured 3-phase progression from formal CFG engine foundations to visual CLI pipelines and benchmarks."
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(11.0)
    p_sub.font.color.rgb = TEXT_MUTED

    # 4. Phase Cards
    card_w = 3450000
    card_h = 4750000
    card_y = 1620000
    gap = 370000

    phases = [
        {
            "num": "PHASE 1",
            "title": "Core Engine & Foundation",
            "status": "COMPLETED (100%)",
            "status_done": True,
            "color": HEADER_ACCENT_1,
            "timeline": "Week 1",
            "objective": "Formal grammar specification, lexical tokenizer, and LL(1) parsing core.",
            "modules": [
                ("grammar/email.cfg", "CFG production rules"),
                ("lexer.py", "Position-aware tokenizer"),
                ("parser.py", "Recursive-descent LL(1)"),
            ],
            "metric_label": "Test Suite Accuracy",
            "metric_val": "10/10 Passed (100%)"
        },
        {
            "num": "PHASE 2",
            "title": "Validation & CLI Pipeline",
            "status": "COMPLETED (100%)",
            "status_done": True,
            "color": HEADER_ACCENT_2,
            "timeline": "Weeks 2–3",
            "objective": "Verdict engine, structured diagnostics, and multi-format visual tree rendering.",
            "modules": [
                ("validator.py", "ACCEPT/REJECT verdict & CLI"),
                ("visualize.py", "SVG, DOT & ASCII trees"),
                ("CLI Pipeline", "Single, batch & REPL mode"),
            ],
            "metric_label": "Test Suite Accuracy",
            "metric_val": "19/19 Passed (100%)"
        },
        {
            "num": "PHASE 3",
            "title": "QA, Benchmarks & Release",
            "status": "COMPLETED (100%)",
            "status_done": True,
            "color": HEADER_ACCENT_3,
            "timeline": "Week 4",
            "objective": "Performance benchmarking, secondary grammar proof, and final documentation.",
            "modules": [
                ("benchmark.py", "O(n) linear latency test"),
                ("grammar/date.cfg", "Grammar extensibility demo"),
                ("tests/ & Docs", "Full QA suite & report"),
            ],
            "metric_label": "Accuracy & Complexity",
            "metric_val": "100% Passed • O(n) Verified"
        }
    ]

    for i, pdata in enumerate(phases):
        cx = 548640 + i * (card_w + gap)

        # Card Base (White rectangle with border)
        card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, card_h)
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = CARD_BG
        card_shape.line.color.rgb = pdata["color"] if pdata["status_done"] else CARD_BORDER
        card_shape.line.width = Pt(1.5 if pdata["status_done"] else 1.0)

        # Top Accent Strip
        top_strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, card_y, card_w, 90000)
        top_strip.fill.solid()
        top_strip.fill.fore_color.rgb = pdata["color"]
        top_strip.line.fill.background()

        # Inner Content Textbox
        pad_x = 220000
        pad_y = 200000
        content_box = slide.shapes.add_textbox(cx + pad_x, card_y + pad_y, card_w - (2 * pad_x), card_h - (2 * pad_y))
        tf = content_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Paragraph 1: Phase Number & Status
        p1 = tf.paragraphs[0]
        r1_num = p1.add_run()
        r1_num.text = f"{pdata['num']}  •  {pdata['timeline']}\n"
        r1_num.font.name = "Arial"
        r1_num.font.size = Pt(9.5)
        r1_num.font.bold = True
        r1_num.font.color.rgb = pdata["color"]

        # Status Pill Text
        r1_stat = p1.add_run()
        r1_stat.text = f"[{pdata['status']}]\n"
        r1_stat.font.name = "Arial"
        r1_stat.font.size = Pt(9.0)
        r1_stat.font.bold = True
        r1_stat.font.color.rgb = BADGE_TEXT_DONE if pdata["status_done"] else BADGE_TEXT_NEXT

        # Paragraph 2: Phase Title
        p2 = tf.add_paragraph()
        p2.space_before = Pt(4)
        r2 = p2.add_run()
        r2.text = f"{pdata['title']}\n"
        r2.font.name = "Arial"
        r2.font.size = Pt(14.0)
        r2.font.bold = True
        r2.font.color.rgb = PRIMARY_DARK

        # Paragraph 3: Objective
        p3 = tf.add_paragraph()
        p3.space_before = Pt(2)
        r3 = p3.add_run()
        r3.text = f"{pdata['objective']}\n"
        r3.font.name = "Arial"
        r3.font.size = Pt(9.5)
        r3.font.color.rgb = TEXT_MUTED

        # Paragraph 4: Key Modules Header
        p4 = tf.add_paragraph()
        p4.space_before = Pt(8)
        r4 = p4.add_run()
        r4.text = "DELIVERABLES & MODULES:\n"
        r4.font.name = "Arial"
        r4.font.size = Pt(9.0)
        r4.font.bold = True
        r4.font.color.rgb = PRIMARY_DARK

        # Paragraph 5: Module items
        for mod_name, mod_desc in pdata["modules"]:
            pm = tf.add_paragraph()
            pm.space_before = Pt(2)
            rm_bullet = pm.add_run()
            rm_bullet.text = "▪ "
            rm_bullet.font.name = "Arial"
            rm_bullet.font.size = Pt(8.5)
            rm_bullet.font.color.rgb = pdata["color"]

            rm_code = pm.add_run()
            rm_code.text = f"{mod_name}: "
            rm_code.font.name = "Arial"
            rm_code.font.size = Pt(9.0)
            rm_code.font.bold = True
            rm_code.font.color.rgb = PRIMARY_DARK

            rm_desc = pm.add_run()
            rm_desc.text = mod_desc
            rm_desc.font.name = "Arial"
            rm_desc.font.size = Pt(8.5)
            rm_desc.font.color.rgb = TEXT_MUTED

        # Paragraph 6: Outcome / Metric Box at bottom
        p_met = tf.add_paragraph()
        p_met.space_before = Pt(14)
        r_met_lbl = p_met.add_run()
        r_met_lbl.text = f"{pdata['metric_label']}\n"
        r_met_lbl.font.name = "Arial"
        r_met_lbl.font.size = Pt(8.5)
        r_met_lbl.font.color.rgb = TEXT_MUTED

        r_met_val = p_met.add_run()
        r_met_val.text = pdata["metric_val"]
        r_met_val.font.name = "Arial"
        r_met_val.font.size = Pt(12.0)
        r_met_val.font.bold = True
        r_met_val.font.color.rgb = pdata["color"]

    # 5. Footer Label (Bottom Right)
    footer_box = slide.shapes.add_textbox(7616952, 6537960, 4114800, 274320)
    tf_foot = footer_box.text_frame
    tf_foot.word_wrap = True
    tf_foot.margin_left = tf_foot.margin_top = tf_foot.margin_right = tf_foot.margin_bottom = 0
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = "GRAMMAR-BASED PATTERN RECOGNITION"
    p_foot.font.name = "Arial"
    p_foot.font.size = Pt(8.0)
    p_foot.font.color.rgb = TEXT_MUTED

    # Save presentation
    prs.save(pptx_path)
    print(f"Successfully updated presentation: '{pptx_path}'")
    print(f"Total slides now: {len(prs.slides)}")


if __name__ == "__main__":
    ppt_file = "Grammar-Based_Pattern_Recognition_TAE1.pptx"
    add_phase_completion_slide(ppt_file)
